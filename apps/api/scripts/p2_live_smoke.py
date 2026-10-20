"""Run a disposable P2 smoke test against the live API and Celery worker."""
from __future__ import annotations

import os
import subprocess
import tempfile
import time
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import httpx
from docx import Document

BASE_URL = os.getenv("P2_SMOKE_BASE_URL", "http://127.0.0.1:9000/api/v1")
TERMINAL_STATUSES = {
    "READY",
    "PARTIAL_SUCCESS",
    "WAITING_FOR_CONFIGURATION",
    "FAILED",
}


def wait_for_api(client: httpx.Client) -> None:
    for _ in range(30):
        try:
            if client.get(f"{BASE_URL}/health/ready").status_code == 200:
                return
        except httpx.HTTPError:
            pass
        time.sleep(1)
    raise RuntimeError("API did not become ready")


def wait_for_asset(client: httpx.Client, asset_id: str) -> dict:
    for _ in range(60):
        response = client.get(f"{BASE_URL}/assets/{asset_id}")
        response.raise_for_status()
        payload = response.json()
        if payload["status"] in TERMINAL_STATUSES:
            return payload
        time.sleep(0.5)
    raise RuntimeError(f"asset {asset_id} did not reach a terminal state")


def make_docx() -> bytes:
    document = Document()
    document.add_heading("P2 真实队列检查", level=1)
    document.add_paragraph("Python 中生成器和迭代器有什么区别？")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_legacy_doc() -> bytes:
    with tempfile.TemporaryDirectory(prefix="bagu-live-doc-") as temp_directory:
        source = Path(temp_directory) / "legacy-source.docx"
        source.write_bytes(make_docx())
        completed = subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "doc:MS Word 97",
                "--outdir",
                temp_directory,
                str(source),
            ],
            env={"HOME": temp_directory, "PATH": os.environ.get("PATH", ""), "LANG": "C.UTF-8"},
            capture_output=True,
            check=False,
            timeout=60,
        )
        result = Path(temp_directory) / "legacy-source.doc"
        if completed.returncode != 0 or not result.exists():
            raise RuntimeError("failed to create disposable legacy DOC fixture")
        return result.read_bytes()


def main() -> None:
    asset_ids: list[str] = []
    with httpx.Client(timeout=10) as client:
        wait_for_api(client)
        collection = client.post(
            f"{BASE_URL}/collections",
            headers={"Idempotency-Key": str(uuid4())},
            json={
                "name": f"P2 live smoke {uuid4()}",
                "purpose": "INTERVIEW_EXPERIENCE",
            },
        )
        collection.raise_for_status()
        collection_id = collection.json()["id"]
        try:
            uploaded = client.post(
                f"{BASE_URL}/collections/{collection_id}/assets/upload",
                headers={"Idempotency-Key": str(uuid4())},
                files={
                    "file": (
                        "p2-smoke.docx",
                        make_docx(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
            uploaded.raise_for_status()
            asset_ids.append(uploaded.json()["id"])
            docx_result = wait_for_asset(client, asset_ids[-1])
            assert docx_result["status"] == "READY", docx_result["job"]
            assert "生成器和迭代器" in docx_result["content"]
            print("DOCX queue: READY")

            uploaded_doc = client.post(
                f"{BASE_URL}/collections/{collection_id}/assets/upload",
                headers={"Idempotency-Key": str(uuid4())},
                files={"file": ("p2-smoke.doc", make_legacy_doc(), "application/msword")},
            )
            uploaded_doc.raise_for_status()
            asset_ids.append(uploaded_doc.json()["id"])
            doc_result = wait_for_asset(client, asset_ids[-1])
            assert doc_result["status"] == "READY", doc_result["job"]
            assert "生成器和迭代器" in doc_result["content"]
            print("Legacy DOC queue: READY")

            submitted_url = client.post(
                f"{BASE_URL}/collections/{collection_id}/assets/url",
                headers={"Idempotency-Key": str(uuid4())},
                json={"url": "https://example.com/", "title": "P2 public web smoke"},
            )
            submitted_url.raise_for_status()
            asset_ids.append(submitted_url.json()["id"])
            web_result = wait_for_asset(client, asset_ids[-1])
            assert web_result["status"] == "READY", web_result["job"]
            assert web_result["sourceUrl"] == "https://example.com/"
            assert web_result["fetchedAt"]
            assert web_result["content"].strip()
            print("Public web queue: READY")

            if not os.getenv("GLM_API_KEY"):
                uploaded = client.post(
                    f"{BASE_URL}/collections/{collection_id}/assets/upload",
                    headers={"Idempotency-Key": str(uuid4())},
                    files={"file": ("p2-smoke.png", b"\x89PNG\r\n\x1a\nsmoke", "image/png")},
                )
                uploaded.raise_for_status()
                asset_ids.append(uploaded.json()["id"])
                image_result = wait_for_asset(client, asset_ids[-1])
                assert image_result["status"] == "WAITING_FOR_CONFIGURATION", image_result["job"]
                print("Image without GLM_API_KEY: WAITING_FOR_CONFIGURATION")
            else:
                print("Image missing-key check: SKIPPED because GLM_API_KEY is configured")
        finally:
            for asset_id in asset_ids:
                client.delete(f"{BASE_URL}/assets/{asset_id}")
            client.delete(f"{BASE_URL}/collections/{collection_id}")


if __name__ == "__main__":
    main()
