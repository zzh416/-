import base64
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from app.models import AssetVersion
from app.modules.content_versions import build_chunks
from app.modules.parsers import IngestionParseError, extract_webpage, parse_docx, parse_pdf


def test_text_chunking_keeps_final_block_before_trailing_newline() -> None:
    content = "第一段\n\n### 最后一道问题？\n"
    version = AssetVersion(version=1, content=content)

    chunks = build_chunks(content, version)

    assert [chunk.content for chunk in chunks] == ["第一段", "### 最后一道问题？"]


def _text_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(output)


def _text_and_image_pdf(text: str) -> bytes:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = (
        f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET\n"
        "q 120 0 0 80 72 560 cm /Im1 Do Q"
    ).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> /XObject << /Im1 6 0 R >> >> "
            b"/Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
        (
            b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 "
            b"/ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 3 >>\n"
            b"stream\n\xff\xff\xff\nendstream"
        ),
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    return bytes(output)


def test_pdf_parser_extracts_text_and_marks_blank_pages_for_ocr() -> None:
    extracted = parse_pdf(
        _text_pdf("Python interview question"),
        max_pages=100,
        max_content_stream_bytes=1024 * 1024,
    )
    assert "Python interview question" in extracted.content
    assert extracted.total_units == 1
    assert extracted.missing_units == []
    assert extracted.blocks[0].block_type == "PDF_PAGE"
    assert extracted.blocks[0].locator == {"page": 1}
    assert extracted.ocr_units == []

    buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(buffer)
    blank = parse_pdf(
        buffer.getvalue(),
        max_pages=100,
        max_content_stream_bytes=1024 * 1024,
    )
    assert blank.total_units == 1
    assert blank.missing_units == [1]
    assert blank.blocks[0].locator == {"page": 1}
    assert blank.ocr_units[0].block_index == 0
    assert blank.ocr_units[0].content_type == "application/pdf"


@pytest.mark.parametrize("text", ["Agent.", "RAG", "A"])
def test_pdf_parser_preserves_short_native_text_without_ocr(text: str) -> None:
    result = parse_pdf(
        _text_pdf(text),
        max_pages=100,
        max_content_stream_bytes=1024 * 1024,
    )

    assert result.blocks[0].content == text
    assert result.missing_units == []
    assert result.ocr_units == []


@pytest.mark.parametrize("text", ["Native text before screenshot", "Agent."])
def test_pdf_parser_marks_text_page_with_embedded_image_for_page_ocr(text: str) -> None:
    result = parse_pdf(
        _text_and_image_pdf(text),
        max_pages=100,
        max_content_stream_bytes=1024 * 1024,
    )

    assert result.blocks[0].content == text
    assert result.blocks[0].locator == {"page": 1}
    assert result.missing_units == []
    assert len(result.ocr_units) == 1
    assert result.ocr_units[0].block_index == 0
    assert result.ocr_units[0].locator == {"page": 1, "reason": "EMBEDDED_IMAGE"}


def test_pdf_parser_rejects_page_count_over_limit() -> None:
    buffer = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.add_blank_page(width=612, height=792)
    writer.write(buffer)

    with pytest.raises(IngestionParseError) as error:
        parse_pdf(
            buffer.getvalue(),
            max_pages=1,
            max_content_stream_bytes=1024 * 1024,
        )
    assert error.value.code == "PDF_TOO_MANY_PAGES"


def test_docx_parser_preserves_paragraph_and_table_order() -> None:
    document = Document()
    document.add_heading("Python 面经", level=1)
    document.add_paragraph("解释 GIL 的作用。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "问题"
    table.cell(0, 1).text = "答案"
    table.cell(1, 0).text = "什么是协程"
    table.cell(1, 1).text = "用户态调度"
    buffer = BytesIO()
    document.save(buffer)

    result = parse_docx(buffer.getvalue())
    assert result.content.index("Python 面经") < result.content.index("解释 GIL")
    assert result.content.index("解释 GIL") < result.content.index("什么是协程")
    assert "| 问题 | 答案 |" in result.content
    assert result.ocr_units == []


def test_docx_parser_keeps_inline_image_between_surrounding_paragraphs() -> None:
    image = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
    )
    document = Document()
    image_paragraph = document.add_paragraph()
    image_paragraph.add_run("图片前的问题")
    image_paragraph.add_run().add_picture(BytesIO(image))
    image_paragraph.add_run("图片后的答案")
    buffer = BytesIO()
    document.save(buffer)

    result = parse_docx(buffer.getvalue())

    assert [block.block_type for block in result.blocks] == [
        "DOCX_PARAGRAPH",
        "DOCX_IMAGE",
        "DOCX_PARAGRAPH",
    ]
    assert result.blocks[0].content == "图片前的问题"
    assert result.blocks[2].content == "图片后的答案"
    assert result.blocks[1].locator["paragraph"] == 1
    assert result.ocr_units[0].block_index == 1
    assert result.ocr_units[0].content == image


def test_webpage_extractor_returns_title_and_main_markdown() -> None:
    html = b"""
    <!doctype html><html><head><title>Python Interview Notes</title></head><body>
      <nav>Home Products Login Contact</nav>
      <article>
        <h1>Python Interview Notes</h1>
        <p>The global interpreter lock protects access to Python objects.</p>
        <p>This article explains threads, processes, and practical tradeoffs.</p>
        <h2>Questions</h2>
        <p>How does the GIL affect CPU-bound workloads?</p>
      </article>
      <footer>Copyright Example</footer>
    </body></html>
    """
    result = extract_webpage(html, "https://example.com/python")
    assert result.title == "Python Interview Notes"
    assert "global interpreter lock" in result.content
    assert "CPU-bound workloads" in result.content
    assert "Products Login Contact" not in result.content


def test_webpage_extractor_rejects_javascript_shell_message() -> None:
    html = b"""
    <!doctype html><html><head><title>SPA</title></head><body>
      <noscript>You need to enable JavaScript to run this app.</noscript>
      <div id="root"></div><script src="/assets/app.js"></script>
    </body></html>
    """

    with pytest.raises(IngestionParseError) as error:
        extract_webpage(html, "https://example.com/app")

    assert error.value.code == "WEB_CONTENT_EMPTY"
