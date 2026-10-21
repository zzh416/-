from io import BytesIO
from unittest.mock import patch
from uuid import UUID, uuid4

from docx import Document
from fastapi.testclient import TestClient

from app.adapters.glm_ocr import OCRNotConfigured, OCRResult
from app.adapters.browser_renderer import RenderedPage
from app.main import app
from app.modules.safe_web import FetchedDocument, SafeWebError
from app.modules.ingestion_service import process_ingestion_job


client = TestClient(app)


def _create_collection() -> str:
    response = client.post(
        "/api/v1/collections",
        json={
            "name": f"P2 Worker 测试-{uuid4()}",
            "purpose": "INTERVIEW_EXPERIENCE",
        },
    )
    assert response.status_code == 201
    return response.json()["id"]


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("FastAPI 面经", level=1)
    document.add_paragraph("依赖注入如何工作？")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_worker_parses_docx_and_persists_version_and_chunks() -> None:
    collection_id = _create_collection()
    asset_id: str | None = None
    try:
        with patch("app.workers.celery_app.celery_app.send_task"):
            uploaded = client.post(
                f"/api/v1/collections/{collection_id}/assets/upload",
                files={
                    "file": (
                        "fastapi.docx",
                        _docx_bytes(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    )
                },
            )
        assert uploaded.status_code == 201
        asset_id = uploaded.json()["id"]
        job_id = uploaded.json()["job"]["id"]

        process_ingestion_job(UUID(job_id))

        detail = client.get(f"/api/v1/assets/{asset_id}")
        assert detail.status_code == 200
        payload = detail.json()
        assert payload["status"] == "READY"
        assert payload["job"]["status"] == "SUCCEEDED"
        assert payload["job"]["stage"] == "COMPLETE"
        assert payload["currentVersion"] == 1
        assert "FastAPI 面经" in payload["content"]
        assert any("依赖注入如何工作" in chunk["content"] for chunk in payload["chunks"])
    finally:
        if asset_id is not None:
            client.delete(f"/api/v1/assets/{asset_id}")
        client.delete(f"/api/v1/collections/{collection_id}")


def test_worker_uses_ocr_for_image_and_waits_when_key_is_missing() -> None:
    collection_id = _create_collection()
    asset_ids: list[str] = []
    image = b"\x89PNG\r\n\x1a\nimage-data"
    try:
        with patch("app.workers.celery_app.celery_app.send_task"):
            first = client.post(
                f"/api/v1/collections/{collection_id}/assets/upload",
                files={"file": ("first.png", image, "image/png")},
            )
            second = client.post(
                f"/api/v1/collections/{collection_id}/assets/upload",
                files={"file": ("second.png", image, "image/png")},
            )
        asset_ids.extend([first.json()["id"], second.json()["id"]])

        process_ingestion_job(
            UUID(first.json()["job"]["id"]),
            ocr_extractor=lambda *args, **kwargs: OCRResult(
                markdown="## 问题\n\nPython 的 GIL 是什么？",
                page_count=1,
                provider_task_id="provider-task",
            ),
        )
        ready = client.get(f"/api/v1/assets/{asset_ids[0]}").json()
        assert ready["status"] == "READY"
        assert "GIL" in ready["content"]

        def no_key(*args, **kwargs):
            raise OCRNotConfigured("missing")

        process_ingestion_job(UUID(second.json()["job"]["id"]), ocr_extractor=no_key)
        waiting = client.get(f"/api/v1/assets/{asset_ids[1]}").json()
        assert waiting["status"] == "WAITING_FOR_CONFIGURATION"
        assert waiting["job"]["errorCode"] == "OCR_NOT_CONFIGURED"
        assert waiting["currentVersion"] == 0

        with patch("app.workers.celery_app.celery_app.send_task") as send_task:
            retried = client.post(f"/api/v1/assets/{asset_ids[1]}/retry")
        assert retried.status_code == 200
        assert retried.json()["status"] == "QUEUED"
        assert retried.json()["job"]["retryCount"] == 1
        send_task.assert_called_once()
    finally:
        for asset_id in asset_ids:
            client.delete(f"/api/v1/assets/{asset_id}")
        client.delete(f"/api/v1/collections/{collection_id}")


def test_worker_fetches_public_webpage_and_persists_snapshot_metadata() -> None:
    collection_id = _create_collection()
    asset_id: str | None = None
    try:
        with patch("app.workers.celery_app.celery_app.send_task"):
            created = client.post(
                f"/api/v1/collections/{collection_id}/assets/url",
                json={"url": "https://example.com/interview", "title": "网页面经"},
            )
        asset_id = created.json()["id"]
        job_id = created.json()["job"]["id"]
        html = b"""
        <html><head><title>Example Interview</title></head><body><article>
        <h1>Python Interview</h1>
        <p>Explain the difference between processes and threads.</p>
        <p>Describe when asynchronous IO improves throughput.</p>
        </article></body></html>
        """

        process_ingestion_job(
            UUID(job_id),
            web_fetcher=lambda *args, **kwargs: FetchedDocument(
                url="https://example.com/interview",
                content_type="text/html",
                content=html,
            ),
        )

        detail = client.get(f"/api/v1/assets/{asset_id}").json()
        assert detail["status"] == "READY"
        assert detail["sourceUrl"] == "https://example.com/interview"
        assert detail["byteSize"] == len(html)
        assert detail["fetchedAt"] is not None
        assert "processes and threads" in detail["content"]
    finally:
        if asset_id is not None:
            client.delete(f"/api/v1/assets/{asset_id}")
        client.delete(f"/api/v1/collections/{collection_id}")


def test_worker_falls_back_to_browser_for_javascript_page() -> None:
    collection_id = _create_collection()
    asset_id: str | None = None
    browser_calls: list[str] = []
    shell = b"""
    <html><head><title>Dynamic shell</title></head><body>
      <div id="app"></div><script src="/assets/app.js"></script>
    </body></html>
    """
    rendered = b"""
    <html><head><title>Rendered Interview Notes</title></head><body><article>
      <h1>Python async interview</h1>
      <p>An event loop schedules coroutines and resumes them when IO is ready.</p>
      <p>This rendered paragraph only exists after JavaScript execution.</p>
    </article></body></html>
    """

    try:
        with patch("app.workers.celery_app.celery_app.send_task"):
            created = client.post(
                f"/api/v1/collections/{collection_id}/assets/url",
                json={"url": "https://example.com/dynamic"},
            )
        asset_id = created.json()["id"]

        def browser_renderer(url: str, **kwargs) -> RenderedPage:
            browser_calls.append(url)
            return RenderedPage(
                url=url,
                content_type="text/html",
                content=rendered,
                title="Rendered Interview Notes",
            )

        process_ingestion_job(
            UUID(created.json()["job"]["id"]),
            web_fetcher=lambda *args, **kwargs: FetchedDocument(
                url="https://example.com/dynamic",
                content_type="text/html",
                content=shell,
            ),
            browser_renderer=browser_renderer,
        )

        detail = client.get(f"/api/v1/assets/{asset_id}").json()
        assert browser_calls == ["https://example.com/dynamic"]
        assert detail["status"] == "READY"
        assert detail["title"] == "Rendered Interview Notes"
        assert detail["byteSize"] == len(rendered)
        assert "event loop schedules coroutines" in detail["content"]
    finally:
        if asset_id is not None:
            client.delete(f"/api/v1/assets/{asset_id}")
        client.delete(f"/api/v1/collections/{collection_id}")


def test_worker_falls_back_to_browser_when_static_fetch_redirects() -> None:
    collection_id = _create_collection()
    asset_id: str | None = None
    rendered = b"""
    <html><head><title>Canonical Article</title></head><body><article>
      <h1>Python descriptors</h1>
      <p>A descriptor defines get, set, or delete behavior for an attribute.</p>
      <p>Functions implement the descriptor protocol to create bound methods.</p>
    </article></body></html>
    """

    try:
        with patch("app.workers.celery_app.celery_app.send_task"):
            created = client.post(
                f"/api/v1/collections/{collection_id}/assets/url",
                json={"url": "https://example.com/redirecting-article"},
            )
        asset_id = created.json()["id"]

        def redirected(*args, **kwargs):
            raise SafeWebError("WEB_REDIRECT_BLOCKED", "redirect")

        process_ingestion_job(
            UUID(created.json()["job"]["id"]),
            web_fetcher=redirected,
            browser_renderer=lambda url, **kwargs: RenderedPage(
                url="https://www.example.com/canonical-article",
                content_type="text/html",
                content=rendered,
                title="Canonical Article",
            ),
        )

        detail = client.get(f"/api/v1/assets/{asset_id}").json()
        assert detail["status"] == "READY"
        assert detail["title"] == "Canonical Article"
        assert "descriptor protocol" in detail["content"]
    finally:
        if asset_id is not None:
            client.delete(f"/api/v1/assets/{asset_id}")
        client.delete(f"/api/v1/collections/{collection_id}")
